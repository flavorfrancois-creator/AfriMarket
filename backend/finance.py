"""Module 4 — Finance helpers: invoice PDF, accounting CSV/XLSX export."""
import io
import csv
from datetime import datetime, timezone

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
import openpyxl
from openpyxl.styles import Font, PatternFill

BRAND = "AfriMarket"
BRAND_COLOR = colors.HexColor("#C0392B")


def _fmt(amount, currency=""):
    try:
        return f"{float(amount):,.2f} {currency}".strip().replace(",", " ")
    except Exception:
        return f"{amount} {currency}".strip()


def build_invoice_pdf(order: dict, shop: dict) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=18 * mm, bottomMargin=18 * mm,
                            leftMargin=16 * mm, rightMargin=16 * mm, title=f"Facture {order.get('ref')}")
    styles = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=styles["Title"], textColor=BRAND_COLOR, fontSize=22, spaceAfter=2)
    small = ParagraphStyle("small", parent=styles["Normal"], fontSize=8.5, textColor=colors.HexColor("#555555"))
    normal = styles["Normal"]
    bold = ParagraphStyle("bold", parent=styles["Normal"], fontName="Helvetica-Bold")
    el = []

    date = (order.get("created_at") or "")[:10]
    header = Table([[
        Paragraph(f"<b>{BRAND}</b><br/><font size=8>Marketplace multi-boutiques</font>", h1),
        Paragraph(f"<b>FACTURE</b><br/>N° {order.get('ref')}<br/>Date : {date}<br/>Statut : {order.get('status')}", small),
    ]], colWidths=[100 * mm, 70 * mm])
    header.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"),
                                ("ALIGN", (1, 0), (1, 0), "RIGHT")]))
    el += [header, Spacer(1, 8 * mm)]

    private = order.get("private_client") or {}
    buyer_name = private.get("name") if order.get("is_private") and private else order.get("customer_name", "")
    buyer_extra = ""
    if order.get("is_private") and private:
        buyer_extra = f"{private.get('phone', '')}<br/>{private.get('address', '')}"
    seller = Paragraph(
        f"<b>Vendeur</b><br/>{shop.get('name', '')}<br/>{shop.get('city', '')}, {shop.get('country', '')}"
        f"<br/>{shop.get('phone', '')}", small)
    buyer = Paragraph(
        f"<b>Client</b><br/>{buyer_name}<br/>{buyer_extra or (order.get('address') or '')}", small)
    party = Table([[seller, buyer]], colWidths=[85 * mm, 85 * mm])
    party.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP")]))
    el += [party, Spacer(1, 6 * mm)]

    cur = order.get("currency", "")
    rows = [["Article", "Qté", "P.U.", "Total"]]
    for it in order.get("items", []):
        rows.append([it.get("name", ""), str(it.get("qty", 0)),
                     _fmt(it.get("unit_price", 0), cur), _fmt(it.get("line_total", 0), cur)])
    table = Table(rows, colWidths=[95 * mm, 18 * mm, 30 * mm, 27 * mm])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), BRAND_COLOR),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("ALIGN", (1, 0), (-1, -1), "RIGHT"),
        ("ALIGN", (0, 0), (0, -1), "LEFT"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#FBEDEA")]),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#E0C9C4")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    el += [table, Spacer(1, 4 * mm)]

    totals = [["Sous-total", _fmt(order.get("subtotal", 0), cur)],
              ["Livraison", _fmt(order.get("shipping", 0), cur)],
              ["TOTAL", _fmt(order.get("total", 0), cur)]]
    if order.get("wallet_paid"):
        totals.insert(2, ["Payé par solde", _fmt(order.get("wallet_paid", 0), cur)])
    tt = Table(totals, colWidths=[130 * mm, 40 * mm])
    tt.setStyle(TableStyle([
        ("ALIGN", (0, 0), (-1, -1), "RIGHT"),
        ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
        ("TEXTCOLOR", (0, -1), (-1, -1), BRAND_COLOR),
        ("FONTSIZE", (0, 0), (-1, -1), 10),
        ("LINEABOVE", (0, -1), (-1, -1), 0.6, BRAND_COLOR),
        ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    el += [tt, Spacer(1, 10 * mm)]
    el += [Paragraph(f"Paiement : {order.get('payment_provider', '')} — {order.get('payment_status', '')}", small)]
    el += [Spacer(1, 4 * mm), Paragraph(f"Merci pour votre confiance. Facture générée par {BRAND}.", small)]

    doc.build(el)
    return buf.getvalue()


ACCOUNTING_COLUMNS = [
    ("date", "Date"), ("ref", "Référence"), ("shop_name", "Boutique"),
    ("customer_name", "Client"), ("status", "Statut"), ("payment_status", "Paiement"),
    ("currency", "Devise"), ("subtotal", "Sous-total"), ("shipping", "Livraison"),
    ("total", "Total"), ("commission", "Commission"), ("net", "Net vendeur"),
]


def accounting_rows(orders: list) -> list:
    rows = []
    for o in orders:
        sub = float(o.get("subtotal", 0) or 0)
        comm = float(o.get("commission", 0) or 0)
        rows.append({
            "date": (o.get("created_at") or "")[:10],
            "ref": o.get("ref", ""),
            "shop_name": o.get("shop_name", ""),
            "customer_name": o.get("customer_name", ""),
            "status": o.get("status", ""),
            "payment_status": o.get("payment_status", ""),
            "currency": o.get("currency", ""),
            "subtotal": round(sub, 2),
            "shipping": round(float(o.get("shipping", 0) or 0), 2),
            "total": round(float(o.get("total", 0) or 0), 2),
            "commission": round(comm, 2),
            "net": round(sub - comm, 2),
        })
    return rows


def build_csv(rows: list) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf, delimiter=";")
    writer.writerow([label for _, label in ACCOUNTING_COLUMNS])
    for r in rows:
        writer.writerow([r.get(key, "") for key, _ in ACCOUNTING_COLUMNS])
    return ("\ufeff" + buf.getvalue()).encode("utf-8")


def build_xlsx(rows: list) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Comptabilité"
    header_fill = PatternFill("solid", fgColor="C0392B")
    for col, (_, label) in enumerate(ACCOUNTING_COLUMNS, start=1):
        c = ws.cell(row=1, column=col, value=label)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = header_fill
    for i, r in enumerate(rows, start=2):
        for col, (key, _) in enumerate(ACCOUNTING_COLUMNS, start=1):
            ws.cell(row=i, column=col, value=r.get(key, ""))
    for col in range(1, len(ACCOUNTING_COLUMNS) + 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 16
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_label(order: dict, shop: dict, carrier: dict) -> bytes:
    from reportlab.graphics.barcode import code128
    from reportlab.graphics.shapes import Drawing
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=16 * mm, bottomMargin=16 * mm,
                            leftMargin=16 * mm, rightMargin=16 * mm, title=f"Etiquette {order.get('ref')}")
    styles = getSampleStyleSheet()
    small = ParagraphStyle("small", parent=styles["Normal"], fontSize=9)
    big = ParagraphStyle("big", parent=styles["Title"], fontSize=20, textColor=BRAND_COLOR, spaceAfter=2)
    el = []
    carrier_name = (carrier or {}).get("name", "Transporteur")
    el.append(Paragraph("ÉTIQUETTE D'EXPÉDITION", big))
    el.append(Paragraph(f"<b>Transporteur :</b> {carrier_name} ({(carrier or {}).get('code', '')})", small))
    el.append(Spacer(1, 6 * mm))

    private = order.get("private_client") or {}
    if order.get("is_private") and private:
        rcpt_name, rcpt_phone, rcpt_addr = private.get("name", ""), private.get("phone", ""), private.get("address", "")
    else:
        rcpt_name, rcpt_phone, rcpt_addr = order.get("customer_name", ""), "", order.get("address", "")
    party = Table([[
        Paragraph(f"<b>EXPÉDITEUR</b><br/>{shop.get('name', '')}<br/>{shop.get('city', '')}, {shop.get('country', '')}<br/>{shop.get('phone', '')}", small),
        Paragraph(f"<b>DESTINATAIRE</b><br/>{rcpt_name}<br/>{rcpt_phone}<br/>{rcpt_addr}", small),
    ]], colWidths=[85 * mm, 85 * mm])
    party.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"),
                               ("BOX", (0, 0), (-1, -1), 0.6, colors.HexColor("#CCCCCC")),
                               ("INNERGRID", (0, 0), (-1, -1), 0.6, colors.HexColor("#CCCCCC")),
                               ("TOPPADDING", (0, 0), (-1, -1), 8), ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
                               ("LEFTPADDING", (0, 0), (-1, -1), 8)]))
    el += [party, Spacer(1, 8 * mm)]

    tracking = order.get("tracking_number", "")
    el.append(Paragraph(f"<b>Commande :</b> {order.get('ref', '')}", small))
    el.append(Paragraph(f"<b>Suivi :</b> {tracking}", small))
    el.append(Spacer(1, 6 * mm))
    try:
        bc = code128.Code128(tracking, barHeight=22 * mm, barWidth=0.5 * mm)
        d = Drawing(bc.width, 24 * mm)
        d.add(bc)
        el.append(d)
    except Exception:
        el.append(Paragraph(f"<font size=18><b>{tracking}</b></font>", small))
    doc.build(el)
    return buf.getvalue()
